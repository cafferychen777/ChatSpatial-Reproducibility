#!/usr/bin/env python
"""Statistics, supplementary tables and source data from the OSCC ChatSpatial 1.5.7 run (cluster compute node only).

Usage: oscc157_postrun.py WORKDIR REPO [MDECONV_RUNS]
  WORKDIR/runs/tables     tables written by oscc157_collect.py
  REPO                    root of this repository
  MDECONV_RUNS            proportions written by the method-concordance runs (deconvolution_concordance/scripts/
                          mdeconv_run.py); default $CHATSPATIAL_WORK_DIR/mdeconv/runs_final
Steps
  A  Fig. 2 statistics           experiments/case_studies/compute_case_study_stats.py (OSCC part)
                                 -> experiments/case_studies/oscc/results/stats
  B  Supplementary Table 1       scripts/build_supplementary_tables.py write_effect_sizes()
  C  Supplementary Table 7       CARD / RCTD / FlashDeconv proportions of the method-concordance runs
                                 with the TC/LE labels of this run, then mdeconv_analyze.py
  D  Source data                 Fig. 2 sheets, Supp_Table_1, Index; source_data.xlsx and source_data.zip
"""
import importlib.util
import json
import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd

W, P = Path(sys.argv[1]), Path(sys.argv[2])
T = W / "runs" / "tables"
SAMPLES = [f"s{i}" for i in range(1, 13)]
CS = P / "experiments" / "case_studies"
CASE = CS / "oscc" / "results" / "stats"
MDC = CS / "deconvolution_concordance"
MDECONV_RUNS = Path(sys.argv[3]) if len(sys.argv) > 3 else \
    Path(os.environ.get("CHATSPATIAL_WORK_DIR", ".")) / "mdeconv" / "runs_final"


def load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# ---- A: statistics --------------------------------------------------------------------------------
os.environ["OSCC_RUN_DIR"], os.environ["OSCC_OUT_DIR"] = str(T), str(CASE)
ccs = load("ccs", CS / "compute_case_study_stats.py")
ccs.run_oscc_stats()

# ---- B: Supplementary Table 1 ---------------------------------------------------------------------
bst = load("bst", P / "scripts" / "build_supplementary_tables.py")
bst.CASE_DIR, bst.OUT_DIR = CASE, P / "supplementary_tables"
bst.write_effect_sizes()

# ---- C: Supplementary Table 7 with the new TC/LE labels ------------------------------------------
labels = {s: pd.read_csv(T / "proportions" / f"{s}.csv", index_col=0)["tc_le_region"] for s in SAMPLES}
src, dst = MDECONV_RUNS, MDC / "results" / "runs"
dst.mkdir(parents=True, exist_ok=True)
spot_check = []
for f in sorted(src.glob("s*__*.csv")):
    sid, run = f.stem.split("__")
    df = pd.read_csv(f, index_col=0)
    new = labels[sid]
    spot_check.append({"sample": sid, "run": run, "spots_run": len(df), "spots_v157": len(new),
                       "common": len(df.index.intersection(new.index)),
                       "same_spot_set": set(df.index) == set(new.index)})
    if run == "manuscript":  # the run that stood for the figure's own proportions: now the 1.5.7 FlashDeconv run
        df = pd.read_csv(T / "proportions" / f"{sid}.csv", index_col=0)
    else:
        df["tc_le_region"] = new.reindex(df.index).fillna("not in v157").values
    df.to_csv(dst / (f.name + ".gz"))
for f in src.glob("*__runlog.json"):
    shutil.copy(f, dst / f.name)
pd.DataFrame(spot_check).to_csv(MDC / "results" / "spot_check_chatspatial157.csv", index=False)
an = MDC / "results" / "analysis"
subprocess.run([sys.executable, "-W", "ignore", str(MDC / "scripts" / "mdeconv_analyze.py"), str(dst), str(an)], check=True)
cc = pd.read_csv(an / "consensus_counts.csv")
pl = pd.read_csv(an / "patient_level_tests.csv")
order = ["Cancer", "Endothelial", "Macrophage", "Fibroblast", "-Fibroblast", "myocyte", "B cell", "Mast", "Dendritic",
         "T cell"]
disp = {"-Fibroblast": "CAF (Myofibroblast)", "myocyte": "Myocyte", "Mast": "Mast cell"}
rows = []
for ct in order:
    row = {"Cell Type": disp.get(ct, ct)}
    for run, lab in (("card", "CARD"), ("rctd", "RCTD"), ("flashdeconv", "FlashDeconv")):
        c = cc[(cc["run"] == run) & (cc["cell_type"] == ct)]
        p = pl[(pl["run"] == run) & (pl["cell_type"] == ct)]
        if c.empty or p.empty:
            row[f"{lab} LE/TC-enriched samples (of 12)"] = "not estimated"
            row[f"{lab} median log2FC"] = ""
            row[f"{lab} patient-level q (BH)"] = ""
            continue
        row[f"{lab} LE/TC-enriched samples (of 12)"] = f"{int(c['le_enriched'].iloc[0])}/{int(c['tc_enriched'].iloc[0])}"
        row[f"{lab} median log2FC"] = f"{p['median_log2fc'].iloc[0]:.2f}"
        q = p["q"].iloc[0]
        row[f"{lab} patient-level q (BH)"] = "" if pd.isna(q) else f"{q:.3f}"
    rows.append(row)
pd.DataFrame(rows).to_csv(P / "supplementary_tables" / "Supplementary_Table_7_Deconvolution_Method_Concordance.csv",
                          index=False)

# Supplementary_Tables.xlsx: replace Table_1 and Table_7
from openpyxl import load_workbook  # noqa: E402

xl = P / "supplementary_tables" / "Supplementary_Tables.xlsx"
if xl.exists():
    wb = load_workbook(xl)
    names = list(wb.sheetnames)
    for sheet, csv in (("Table_1", "Supplementary_Table_1_Effect_Sizes.csv"),
                       ("Table_7", "Supplementary_Table_7_Deconvolution_Method_Concordance.csv")):
        pos = names.index(sheet)
        del wb[sheet]
        ws = wb.create_sheet(sheet, pos)
        df = pd.read_csv(P / "supplementary_tables" / csv, keep_default_na=False)
        ws.append(list(df.columns))
        for r in df.itertuples(index=False):
            ws.append([None if v == "" else v for v in r])
    wb.save(xl)

# ---- D: source data -------------------------------------------------------------------------------
SD = P / "source_data"
spatial_rows = []
for sid in ("s1", "s9"):
    gsm = f"GSM{6339630 + int(sid[1:])}"
    vdir = W / "input" / "visium" / sid / "spatial"
    pos = pd.read_csv(vdir / "tissue_positions_list.csv", header=None,
                      names=["barcode", "in_tissue", "row", "col", "pixel_row", "pixel_col"])
    pos = pos[pos["in_tissue"] == 1].copy()
    scale = json.load(open(vdir / "scalefactors_json.json")).get("tissue_hires_scalef", 1.0)
    pos["pixel_row"] *= scale
    pos["pixel_col"] *= scale
    spots = pd.read_csv(T / "spots" / f"{sid}.csv", index_col=0)
    props = pd.read_csv(T / "proportions" / f"{sid}.csv", index_col=0)
    out = pos[["barcode", "pixel_row", "pixel_col"]].copy()
    out.insert(0, "sample", sid)
    out["tc_le_region"] = spots["tc_le_region"].reindex(out["barcode"]).values
    out["fibroblast_proportion"] = props["Fibroblast"].reindex(out["barcode"]).values
    out["FN1_expression"] = spots["FN1_lognorm"].reindex(out["barcode"]).values
    spatial_rows.append(out)
sheets = {
    "Fig2_spatial_maps": (pd.concat(spatial_rows, ignore_index=True),
                          "ChatSpatial 1.5.7 run: tissue positions, TC/LE labels, FlashDeconv proportions, FN1"),
    "Fig2d_enrichment": (pd.read_csv(CASE / "per_sample_tc_le_enrichment.csv"), "per_sample_tc_le_enrichment.csv"),
    "Fig2e_consensus": (pd.read_csv(CASE / "consensus_tc_le_patterns.csv"), "consensus_tc_le_patterns.csv"),
    "Fig2e_consensus_CI": (pd.read_csv(CASE / "consensus_binomial_ci.csv"), "consensus_binomial_ci.csv"),
    "Fig2f_cellchat": (pd.read_csv(T / "cellchat_tc_le_caf_cancer_interactions.csv"),
                       "cellchat_tc_le_caf_cancer_interactions.csv (analyze_cell_communication, cellchat_r)"),
    "Fig2h_morans": (pd.read_csv(T / "moranI_results_all_samples.csv"),
                     "moranI_results_all_samples.csv (analyze_spatial_statistics, moran)"),
    "Fig2i_DE_all_genes": (pd.read_csv(T / "deg_tc_vs_le_top500.csv"),
                           "deg_tc_vs_le_top500.csv (find_markers, top 500 genes per direction)"),
}
mk = []
for g in ("tc", "le"):
    d = pd.read_csv(CASE / f"consensus_{g}_upregulated_genes.csv").rename(columns={"sample": "n_samples"})
    d["region"] = g.upper()
    d["source_category"] = f"consensus_{g}_upregulated"
    mk.append(d)
svg = pd.read_csv(CASE / "svg_target_genes_summary.csv").rename(columns={"category": "region"})
svg["source_category"] = "moranI_sample1"
mk.append(svg)
sheets["Fig2i_marker_genes"] = (pd.concat(mk, ignore_index=True, sort=False),
                                "consensus gene lists and Sample 1 Moran's I summary")
sheets["Supp_Table_1"] = (pd.read_csv(P / "supplementary_tables" / "Supplementary_Table_1_Effect_Sizes.csv"),
                          "supplementary_tables/Supplementary_Table_1_Effect_Sizes.csv")
idx = pd.read_csv(SD / "Index.csv")
for name, (df, source) in sheets.items():
    df.to_csv(SD / f"{name}.csv", index=False)
    m = idx["sheet"] == name
    idx.loc[m, "rows"] = len(df)
    idx.loc[m, "columns"] = df.shape[1]
    idx.loc[m, "source"] = source
idx.loc[idx["sheet"] == "Fig2h_morans", "description"] = "Figure 2h Moran's I values (all sections; panel shows Sample 1)."
idx.loc[idx["sheet"] == "Fig2i_DE_all_genes", "description"] = "TC vs LE marker genes underlying the consensus lists."
idx.to_csv(SD / "Index.csv", index=False)
with pd.ExcelWriter(P / "source_data.xlsx", engine="openpyxl") as w:
    for name in idx["sheet"]:
        pd.read_csv(SD / f"{name}.csv", low_memory=False).to_excel(w, sheet_name=name, index=False)
    idx.to_excel(w, sheet_name="Index", index=False)
with zipfile.ZipFile(P / "source_data.zip", "w", zipfile.ZIP_DEFLATED) as z:
    z.writestr("source_data/", "")
    for f in sorted(SD.glob("*.csv")):
        z.write(f, f"source_data/{f.name}")
print(pd.DataFrame(spot_check).groupby("run")["same_spot_set"].all())
print(open(an / "decision.txt").read())
print("POSTRUN_DONE")
